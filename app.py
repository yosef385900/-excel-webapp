import os, secrets, shutil
from functools import wraps
from flask import Flask, render_template, request, jsonify, session, redirect, url_for, send_file
from openpyxl import load_workbook

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', secrets.token_hex(32))
PASSWORD = os.environ.get('SITE_PASSWORD', '123456')
FILE = os.path.join(os.path.dirname(__file__), 'data.xlsx')


def auth_required(fn):

    @wraps(fn)

    def wrapper(*args, **kwargs):

        if not session.get('ok'):

            if request.path.startswith('/api/'):

                return jsonify({'error': 'unauthorized'}), 401

            return redirect(url_for('login'))

        return fn(*args, **kwargs)

    return wrapper

@app.route('/login', methods=['GET','POST'])
def login():
    error = None
    if request.method == 'POST':
        if secrets.compare_digest(request.form.get('password',''), PASSWORD):
            session['ok'] = True
            return redirect(url_for('index'))
        error = 'كلمة السر غير صحيحة'
    return render_template('login.html', error=error)

@app.route('/logout')
def logout():
    session.clear(); return redirect(url_for('login'))

@app.route('/')
@auth_required
def index():
    return render_template('index.html')

@app.route('/api/sheets')
@auth_required
def sheets():
    wb = load_workbook(FILE, data_only=False, read_only=True)
    return jsonify(wb.sheetnames)

@app.route('/api/sheet/<int:idx>')
@auth_required
def sheet(idx):
    wb = load_workbook(FILE, data_only=False, read_only=True)
    if idx < 0 or idx >= len(wb.worksheets): return jsonify({'error':'bad sheet'}), 404
    ws = wb.worksheets[idx]
    data=[]
    for row in ws.iter_rows():
        data.append([c.value for c in row])
    return jsonify({'name':ws.title,'rows':data,'max_row':ws.max_row,'max_col':ws.max_column})

@app.route('/api/cell', methods=['POST'])
@auth_required
def cell():
    p=request.get_json(force=True)
    idx=int(p['sheet']); row=int(p['row']); col=int(p['col']); value=p.get('value','')
    wb=load_workbook(FILE, data_only=False)
    ws=wb.worksheets[idx]
    if value == '': value = None
    elif isinstance(value,str) and value.startswith('='): pass
    ws.cell(row=row,column=col).value=value
    tmp=FILE+'.tmp.xlsx'; wb.save(tmp); os.replace(tmp,FILE)
    return jsonify({'ok':True})

@app.route('/download')
@auth_required
def download():
    return send_file(FILE, as_attachment=True, download_name='updated.xlsx')

@app.route('/health')
def health(): return 'ok'

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT','5000')))
